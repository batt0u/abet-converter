from __future__ import annotations

import csv
import io
import sqlite3
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from abet_converter.drivers.mdbtools import MdbtoolsRuntime


SUPPORTED_DB_SUFFIXES = {".abetdb", ".mdb"}
SUPPORTED_FORMATS = ("sqlite", "sql", "csv", "xlsx")
EXCEL_MAX_ROWS = 1_048_576


@dataclass(frozen=True)
class ConversionTarget:
    source_db: Path
    sqlite_path: Path | None
    sql_dump_path: Path | None
    csv_dir: Path | None
    xlsx_path: Path | None


def discover_abet_sources(source: Path, recursive: bool = False) -> list[Path]:
    resolved = source.resolve()
    if resolved.is_file():
        if resolved.suffix.lower() not in SUPPORTED_DB_SUFFIXES:
            raise ValueError(f"Unsupported source extension: {resolved}")
        return [resolved]
    if not resolved.is_dir():
        raise FileNotFoundError(f"Source not found: {resolved}")
    pattern = "**/*" if recursive else "*"
    files = [
        item.resolve()
        for item in resolved.glob(pattern)
        if item.is_file() and item.suffix.lower() in SUPPORTED_DB_SUFFIXES
    ]
    return sorted(files)


def normalize_formats(formats: list[str] | None) -> tuple[str, ...]:
    if not formats:
        return ("sqlite",)
    normalized: list[str] = []
    for item in formats:
        value = item.lower()
        if value not in SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported format: {item}")
        if value not in normalized:
            normalized.append(value)
    return tuple(normalized)


def run_mdbtools(runtime: MdbtoolsRuntime, command: list[str | Path]) -> str:
    completed = subprocess.run(
        [str(item) for item in command],
        text=False,
        capture_output=True,
        check=True,
        env=runtime.subprocess_env(),
    )
    return completed.stdout.decode("latin-1", errors="replace")


def get_tables(runtime: MdbtoolsRuntime, source_db: Path) -> list[str]:
    output = run_mdbtools(runtime, [runtime.executable_path("mdb-tables"), "-1", source_db])
    return [line.strip() for line in output.splitlines() if line.strip()]


def get_schema_sql(runtime: MdbtoolsRuntime, source_db: Path) -> str:
    return run_mdbtools(runtime, [runtime.executable_path("mdb-schema"), source_db, "sqlite"])


def export_table_csv(runtime: MdbtoolsRuntime, source_db: Path, table_name: str) -> str:
    return run_mdbtools(
        runtime,
        [
            runtime.executable_path("mdb-export"),
            "-0",
            "__NULL__",
            "-b",
            "hex",
            source_db,
            table_name,
        ],
    )


def parse_csv_text(csv_text: str) -> tuple[list[str], list[list[str | None]]]:
    reader = csv.reader(io.StringIO(csv_text))
    headers = next(reader, None)
    if not headers:
        return [], []
    rows = [[None if value == "__NULL__" else value for value in row] for row in reader]
    return headers, rows


def import_rows_into_sqlite(conn: sqlite3.Connection, table_name: str, headers: list[str], rows: list[list[str | None]]) -> int:
    if not headers:
        return 0
    columns_sql = ", ".join(f'"{column}"' for column in headers)
    placeholders = ", ".join("?" for _ in headers)
    insert_sql = f'INSERT INTO "{table_name}" ({columns_sql}) VALUES ({placeholders})'
    batch: list[list[str | None]] = []
    count = 0
    for row in rows:
        batch.append(row)
        count += 1
        if len(batch) >= 10000:
            conn.executemany(insert_sql, batch)
            batch.clear()
    if batch:
        conn.executemany(insert_sql, batch)
    return count


def write_table_csv(path: Path, headers: list[str], rows: list[list[str | None]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])


def normalize_sheet_name(table_name: str, used_names: set[str]) -> str:
    candidate = table_name[:31] or "Sheet"
    counter = 1
    while candidate in used_names:
        suffix = f"_{counter}"
        candidate = f"{table_name[: max(1, 31 - len(suffix))]}{suffix}"
        counter += 1
    used_names.add(candidate)
    return candidate


def build_conversion_targets(
    sources: list[Path],
    input_path: Path,
    output_path: Path,
    formats: tuple[str, ...],
) -> list[ConversionTarget]:
    if input_path.is_file():
        source_db = sources[0]
        stem = source_db.stem
        output_is_file = formats == ("sqlite",)
        base_dir = output_path.parent if output_is_file else output_path
        sqlite_path = output_path if "sqlite" in formats and output_is_file else (base_dir / f"{stem}.sqlite" if "sqlite" in formats else None)
        return [
            ConversionTarget(
                source_db=source_db,
                sqlite_path=sqlite_path,
                sql_dump_path=base_dir / f"{stem}.sql" if "sql" in formats else None,
                csv_dir=base_dir / f"{stem}_csv" if "csv" in formats else None,
                xlsx_path=base_dir / f"{stem}.xlsx" if "xlsx" in formats else None,
            )
        ]

    targets: list[ConversionTarget] = []
    for source_db in sources:
        relative_path = source_db.relative_to(input_path)
        relative_parent = relative_path.parent
        stem = source_db.stem
        base_dir = output_path / relative_parent
        targets.append(
            ConversionTarget(
                source_db=source_db,
                sqlite_path=base_dir / f"{stem}.sqlite" if "sqlite" in formats else None,
                sql_dump_path=base_dir / f"{stem}.sql" if "sql" in formats else None,
                csv_dir=base_dir / f"{stem}_csv" if "csv" in formats else None,
                xlsx_path=base_dir / f"{stem}.xlsx" if "xlsx" in formats else None,
            )
        )
    return targets


def ensure_target_paths_do_not_exist(target: ConversionTarget) -> None:
    for path in (target.sqlite_path, target.sql_dump_path, target.xlsx_path):
        if path and path.exists():
            raise FileExistsError(f"Output already exists: {path}")
    if target.csv_dir and target.csv_dir.exists():
        raise FileExistsError(f"Output already exists: {target.csv_dir}")


def create_sqlite_database(sqlite_path: Path, schema_sql: str, exported_tables: list[tuple[str, list[str], list[list[str | None]]]]) -> None:
    sqlite_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(sqlite_path))
    try:
        conn.execute("PRAGMA journal_mode = WAL")
        conn.execute("PRAGMA synchronous = NORMAL")
        conn.executescript(schema_sql)
        for table_name, headers, rows in exported_tables:
            with conn:
                import_rows_into_sqlite(conn, table_name, headers, rows)
    finally:
        conn.close()


def write_sql_dump(sqlite_path: Path, sql_dump_path: Path) -> None:
    sql_dump_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(sqlite_path))
    try:
        with sql_dump_path.open("w", encoding="utf-8") as handle:
            for statement in conn.iterdump():
                handle.write(statement)
                handle.write("\n")
    finally:
        conn.close()


def write_xlsx_workbook(
    xlsx_path: Path,
    exported_tables: list[tuple[str, list[str], list[list[str | None]]]],
    max_rows_per_sheet: int = EXCEL_MAX_ROWS,
) -> None:
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    used_names: set[str] = set()
    data_rows_per_sheet = max_rows_per_sheet - 1
    if data_rows_per_sheet < 1:
        raise ValueError("max_rows_per_sheet must leave room for a header and at least one data row.")

    for table_name, headers, rows in exported_tables:
        chunks = [rows[index : index + data_rows_per_sheet] for index in range(0, len(rows), data_rows_per_sheet)] or [[]]
        for chunk_index, chunk in enumerate(chunks, start=1):
            sheet_base_name = table_name if chunk_index == 1 else f"{table_name}_{chunk_index}"
            worksheet = workbook.create_sheet(title=normalize_sheet_name(sheet_base_name, used_names))
            worksheet.append(headers)
            for row in chunk:
                worksheet.append(row)
    workbook.save(xlsx_path)


def convert_one(target: ConversionTarget, runtime: MdbtoolsRuntime) -> list[Path]:
    ensure_target_paths_do_not_exist(target)
    tables = get_tables(runtime, target.source_db)
    if not tables:
        raise RuntimeError(f"No tables found in source DB: {target.source_db}")

    schema_sql = get_schema_sql(runtime, target.source_db)
    exported_tables: list[tuple[str, list[str], list[list[str | None]]]] = []
    written_paths: list[Path] = []

    for table_name in tables:
        csv_text = export_table_csv(runtime, target.source_db, table_name)
        headers, rows = parse_csv_text(csv_text)
        exported_tables.append((table_name, headers, rows))
        if target.csv_dir:
            csv_path = target.csv_dir / f"{table_name}.csv"
            write_table_csv(csv_path, headers, rows)
            written_paths.append(csv_path)

    sqlite_for_dump: Path | None = None
    temporary_sqlite_path: Path | None = None
    if target.sqlite_path:
        create_sqlite_database(target.sqlite_path, schema_sql, exported_tables)
        sqlite_for_dump = target.sqlite_path
        written_paths.append(target.sqlite_path)
    elif target.sql_dump_path:
        temp_dir = Path(tempfile.mkdtemp(prefix="abet_converter_sql_dump_"))
        temporary_sqlite_path = temp_dir / f"{target.source_db.stem}.sqlite"
        create_sqlite_database(temporary_sqlite_path, schema_sql, exported_tables)
        sqlite_for_dump = temporary_sqlite_path

    if target.sql_dump_path:
        if sqlite_for_dump is None:
            raise RuntimeError("SQL dump generation requires a temporary SQLite database.")
        write_sql_dump(sqlite_for_dump, target.sql_dump_path)
        written_paths.append(target.sql_dump_path)

    if target.xlsx_path:
        write_xlsx_workbook(target.xlsx_path, exported_tables)
        written_paths.append(target.xlsx_path)

    if temporary_sqlite_path:
        for extra in (
            temporary_sqlite_path,
            temporary_sqlite_path.with_suffix(f"{temporary_sqlite_path.suffix}-wal"),
            temporary_sqlite_path.with_suffix(f"{temporary_sqlite_path.suffix}-shm"),
        ):
            if extra.exists():
                extra.unlink()
        temporary_sqlite_path.parent.rmdir()

    return written_paths


def convert_many(targets: list[ConversionTarget], runtime: MdbtoolsRuntime) -> list[Path]:
    written_paths: list[Path] = []
    for target in targets:
        written_paths.extend(convert_one(target, runtime))
    return written_paths
